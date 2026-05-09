import os
import datetime
import logging

logger = logging.getLogger(__name__)

def extract_metadata(file_path: str, original_filename: str) -> dict:
    """
    Extract robust metadata from binary and text files.
    """
    ext = os.path.splitext(original_filename)[1].lower()
    
    metadata = {
        "title": original_filename,
        "author": "Unknown",
        "created": "Unknown",
        "modified": "Unknown",
        "pages": "Unknown",
        "words": "Unknown",
        "format": ext.strip('.') or "Unknown",
        "source": original_filename
    }
    
    try:
        # File stat fallback dates
        stat = os.stat(file_path)
        metadata["created"] = datetime.datetime.fromtimestamp(stat.st_ctime).isoformat()
        metadata["modified"] = datetime.datetime.fromtimestamp(stat.st_mtime).isoformat()
    except Exception:
        pass

    try:
        if ext == '.pdf':
            # Try to use PyPDF2 or pypdf
            try:
                try:
                    from pypdf import PdfReader
                except ImportError:
                    from PyPDF2 import PdfReader
                reader = PdfReader(file_path)
                metadata["pages"] = str(len(reader.pages))
                if reader.metadata:
                    metadata["author"] = reader.metadata.get('/Author', metadata["author"]) or metadata["author"]
                    metadata["title"] = reader.metadata.get('/Title', metadata["title"]) or metadata["title"]
            except ImportError:
                logger.debug("PyPDF2/pypdf not available for PDF metadata extraction")
                pass
                
        elif ext == '.docx':
            try:
                import docx
                doc = docx.Document(file_path)
                props = doc.core_properties
                metadata["author"] = props.author or metadata["author"]
                metadata["title"] = props.title or metadata["title"]
                if props.created:
                    metadata["created"] = props.created.isoformat()
                if props.modified:
                    metadata["modified"] = props.modified.isoformat()
                metadata["pages"] = "N/A (Flow document)"
            except ImportError:
                logger.debug("python-docx not available for DOCX metadata extraction")
                pass
                
        elif ext == '.xlsx':
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                metadata["pages"] = f"{len(wb.sheetnames)} sheets"
                props = getattr(wb, "properties", None)
                if props:
                    metadata["author"] = getattr(props, "creator", metadata["author"]) or metadata["author"]
                    metadata["title"] = getattr(props, "title", metadata["title"]) or metadata["title"]
                    created = getattr(props, "created", None)
                    if created:
                        metadata["created"] = created.isoformat()
            except ImportError:
                logger.debug("openpyxl not available for XLSX metadata extraction")
                pass
                
        elif ext == '.pptx':
            try:
                from pptx import Presentation
                prs = Presentation(file_path)
                metadata["pages"] = f"{len(prs.slides)} slides"
                props = prs.core_properties
                if props:
                    metadata["author"] = props.author or metadata["author"]
                    metadata["title"] = props.title or metadata["title"]
                    if props.created:
                        metadata["created"] = props.created.isoformat()
            except ImportError:
                logger.debug("python-pptx not available for PPTX metadata extraction")
                pass

    except Exception as e:
        logger.warning(f"Error extracting deep metadata for {file_path}: {e}")
        
    return metadata

def format_metadata_header(metadata: dict) -> str:
    """Format the metadata into a YAML-style frontmatter block."""
    return f"""---
title: {metadata.get('title', 'Unknown')}
author: {metadata.get('author', 'Unknown')}
created: {metadata.get('created', 'Unknown')}
modified: {metadata.get('modified', 'Unknown')}
pages: {metadata.get('pages', 'Unknown')}
words: {metadata.get('words', 'Unknown')}
format: {metadata.get('format', 'Unknown')}
source: {metadata.get('source', 'Unknown')}
---

"""
